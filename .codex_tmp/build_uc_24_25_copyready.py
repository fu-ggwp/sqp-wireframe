# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Section 2.4 2.5 Use Case Specifications - Copy Ready.docx"

CREATED_BY = "Nguyễn Khắc Thành Sơn"
DATE_CREATED = "25/5/2026"

FIELDS = [
    "ID and Name",
    "Created By",
    "Date Created",
    "Primary Actor",
    "Secondary Actors",
    "Description",
    "Trigger",
    "Preconditions",
    "Postconditions",
    "Normal Flow",
    "Alternative Flows",
    "Exceptions",
    "Priority",
    "Frequency of Use",
    "Business Rules",
    "Other Information",
    "Assumptions",
]


def set_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [("Heading 1", 16, "1F4D78"), ("Heading 2", 14, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Section 2.4 & 2.5 Use Case Specifications - Copy Ready")
    set_font(r, size=9, color="666666")


def add_label_value(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    lr = p.add_run(f"{label}:")
    set_font(lr, bold=True, color="1F4D78")
    if value:
        lines = value.split("\n")
        first = True
        for line in lines:
            if first:
                p.add_run(" ")
                vr = p.add_run(line)
                set_font(vr)
                first = False
            else:
                p.add_run("\n")
                vr = p.add_run(line)
                set_font(vr)


def add_uc(doc: Document, number: str, title: str, data: dict[str, str]) -> None:
    doc.add_heading(f"{number} {title}", level=3)
    for field in FIELDS:
        add_label_value(doc, field, data.get(field, ""))


def uc(id_name, primary, secondary, desc, trigger, pre, post, normal, alt, exc, freq, br, other, assumptions, priority="High"):
    return {
        "ID and Name": id_name,
        "Created By": CREATED_BY,
        "Date Created": DATE_CREATED,
        "Primary Actor": primary,
        "Secondary Actors": secondary,
        "Description": desc,
        "Trigger": trigger,
        "Preconditions": pre,
        "Postconditions": post,
        "Normal Flow": normal,
        "Alternative Flows": alt,
        "Exceptions": exc,
        "Priority": priority,
        "Frequency of Use": freq,
        "Business Rules": br,
        "Other Information": other,
        "Assumptions": assumptions,
    }


def build_data():
    s24 = [
        ("2.4.1", "View question banks", uc(
            "UC-33 View question banks",
            "Teacher",
            "None",
            "As a Teacher, I want to view the Question Banks that I own or am authorized to access so that I can reuse stored questions for Study Set or Exam creation.",
            "The Teacher clicks the Question Banks menu or opens the Question Bank Repository page.",
            "PRE-1. The Teacher has logged in successfully.\nPRE-2. The Teacher account is active and not locked.\nPRE-3. The active role is Teacher.\nPRE-4. The Teacher has at least one owned or authorized Question Bank, or the system can display an empty repository list.",
            "POST-1. The system displays the Teacher's owned or authorized Question Banks.\nPOST-2. The Teacher can select a Question Bank to manage metadata, author questions, or reuse questions for Study Set or Exam creation.\nPOST-3. Learners and Guests do not gain direct access to any Question Bank.",
            "UC-33.0 View Question Banks\n1. The Teacher opens the Question Banks page.\n2. The system validates that the current active role is Teacher.\n3. The system retrieves Question Banks owned by or authorized for the Teacher.\n4. The system displays repository information such as title, subject, topic, visibility, question count, and last updated time.\n5. The Teacher filters, searches, or sorts the repository list if needed.\n6. The Teacher selects a Question Bank to view its detail page.",
            "UC-33.1 No Question Banks Available\n1. At step 3 of the Normal Flow, the system finds no available Question Bank.\n2. The system displays an empty state and provides the Create Question Bank action.\n3. The Teacher may proceed to UC-34 Create question bank.",
            "UC-33.0.E1 Unauthorized Access\n1. At step 2 of the Normal Flow, the system detects that the user is not using the Teacher role.\n2. The system blocks access and displays MSG11.\nUC-33.0.E2 Load Failure\n1. At step 3 of the Normal Flow, the system cannot load repository data.\n2. The system displays MSG13.",
            "Medium. Teachers use this feature when maintaining reusable questions or preparing Study Sets and Exams.",
            "BR-01, BR-09, BR-16, BR-17, BR-37",
            "Question Banks are internal Teacher repositories and are not searchable through public discovery.",
            "The system can determine Question Bank ownership and authorization scope for each Teacher."
        )),
        ("2.4.2", "Create question bank", uc(
            "UC-34 Create question bank",
            "Teacher",
            "None",
            "As a Teacher, I want to create a Question Bank as a reusable repository so that I can store questions for repeated Study Set or Exam creation.",
            "The Teacher clicks the Create Question Bank action from the Question Banks page.",
            "PRE-1. The Teacher has logged in successfully.\nPRE-2. The Teacher account is active and not locked.\nPRE-3. The active role is Teacher.",
            "POST-1. A new Question Bank repository is created successfully.\nPOST-2. The Question Bank is available for shared Question Management use cases.\nPOST-3. No Learner or Guest receives direct access to the Question Bank.",
            "UC-34.0 Create Question Bank\n1. The Teacher opens the Question Banks page.\n2. The Teacher clicks Create Question Bank.\n3. The system displays the Create Question Bank form.\n4. The Teacher enters title, description, subject, topic, visibility, and repository configuration.\n5. The Teacher submits the form.\n6. The system validates required metadata and Teacher authorization.\n7. The system creates the Question Bank.\n8. The system displays a success message.\n9. The system opens the new Question Bank detail page.\n10. The Teacher may add questions through UC-37, UC-38, or UC-43.",
            "UC-34.1 Cancel Creation\n1. At step 4 of the Normal Flow, the Teacher clicks Cancel.\n2. The system discards unsaved data.\n3. The system returns the Teacher to the Question Banks page.",
            "UC-34.0.E1 Invalid Repository Data\n1. At step 6 of the Normal Flow, the system detects missing or invalid required metadata.\n2. The system highlights invalid fields and displays validation messages.\n3. The Teacher corrects the data.\n4. The flow continues from step 4 of the Normal Flow.\nUC-34.0.E2 Unauthorized Role\n1. At step 6, the system detects that the user is not using the Teacher role.\n2. The system blocks the action and displays MSG11.",
            "Medium. Teachers create repositories when preparing repeated learning or exam materials.",
            "BR-01, BR-09, BR-16, BR-17, BR-18, BR-37",
            "Creating a Question Bank creates repository metadata only. Question content is added through shared Question Management.",
            "The system supports empty Question Banks and allows questions to be added later."
        )),
        ("2.4.3", "Update question bank information", uc(
            "UC-35 Update question bank information",
            "Teacher",
            "None",
            "As a Teacher, I want to update Question Bank information so that the repository metadata remains accurate and useful for future reuse.",
            "The Teacher clicks Edit or Update on a Question Bank detail page.",
            "PRE-1. The Teacher has logged in successfully.\nPRE-2. The Teacher account is active and not locked.\nPRE-3. The active role is Teacher.\nPRE-4. The selected Question Bank exists.\nPRE-5. The Teacher owns or is authorized to update the selected Question Bank.",
            "POST-1. The selected Question Bank metadata is updated successfully.\nPOST-2. Existing questions inside the Question Bank remain unchanged unless the Teacher uses Question Management use cases.\nPOST-3. Learners and Guests do not gain direct access to the Question Bank.",
            "UC-35.0 Update Question Bank Information\n1. The Teacher opens an authorized Question Bank detail page.\n2. The Teacher clicks Edit or Update Question Bank.\n3. The system displays existing repository metadata.\n4. The Teacher updates title, description, subject, topic, visibility, or repository configuration.\n5. The Teacher submits the updated metadata.\n6. The system validates the updated data and Teacher authorization.\n7. The system saves the changes.\n8. The system displays a success message.\n9. The system displays the updated Question Bank detail page.",
            "UC-35.1 Cancel Update\n1. At step 4 of the Normal Flow, the Teacher clicks Cancel.\n2. The system discards unsaved changes.\n3. The system returns the Teacher to the Question Bank detail page.\nUC-35.2 Update Visibility\n1. At step 4, the Teacher changes visibility or repository configuration.\n2. The system applies repository visibility rules without granting Learners direct access.\n3. The flow continues from step 6 of the Normal Flow.",
            "UC-35.0.E1 Invalid Metadata\n1. At step 6, the system detects missing or invalid metadata.\n2. The system highlights invalid fields and displays validation messages.\n3. The Teacher corrects the data.\n4. The flow continues from step 4.\nUC-35.0.E2 Unauthorized Repository\n1. At step 6, the system detects that the Teacher is not authorized to update the selected Question Bank.\n2. The system blocks the update and displays MSG11.",
            "Medium. Teachers use this feature when maintaining reusable repository information.",
            "BR-01, BR-09, BR-16, BR-17, BR-20, BR-37",
            "This use case updates repository metadata only. Question content is updated through UC-41 Update question.",
            "The system can separate repository metadata changes from question content changes."
        )),
        ("2.4.4", "Delete question bank", uc(
            "UC-36 Delete question bank",
            "Teacher",
            "None",
            "As a Teacher, I want to delete, archive, or hide a Question Bank that is no longer needed while preserving historical learning, exam, and reporting integrity.",
            "The Teacher clicks Delete, Archive, or Hide on a Question Bank detail page.",
            "PRE-1. The Teacher has logged in successfully.\nPRE-2. The Teacher account is active and not locked.\nPRE-3. The active role is Teacher.\nPRE-4. The selected Question Bank exists.\nPRE-5. The Teacher owns or is authorized to delete, archive, or hide the selected Question Bank.",
            "POST-1. The Question Bank is deleted, archived, or hidden according to integrity rules.\nPOST-2. Historical exam, learning, and report records remain valid.\nPOST-3. Study Sets created from copied Question Bank questions are not deleted automatically.",
            "UC-36.0 Delete Question Bank\n1. The Teacher opens an authorized Question Bank detail page.\n2. The Teacher clicks Delete, Archive, or Hide.\n3. The system displays a confirmation message.\n4. The Teacher confirms the action.\n5. The system checks ownership, authorization, and references from exams, reports, or copied Study Sets.\n6. The system applies the allowed action according to integrity rules.\n7. The system displays a success message.\n8. The system returns the Teacher to the Question Banks page.",
            "UC-36.1 Cancel Deletion\n1. At step 3, the Teacher cancels the confirmation.\n2. The system closes the confirmation message.\n3. The Question Bank remains unchanged.\nUC-36.2 Archive Instead Of Hard Delete\n1. At step 5, the system detects historical references that must be preserved.\n2. The system archives or hides the Question Bank instead of hard-deleting it.\n3. The flow continues from step 7.",
            "UC-36.0.E1 Unauthorized Action\n1. At step 5, the system detects that the Teacher is not authorized to delete, archive, or hide the Question Bank.\n2. The system blocks the action and displays MSG11.\nUC-36.0.E2 Integrity Constraint Blocks Deletion\n1. At step 5, the system detects references that do not allow deletion.\n2. The system prevents hard deletion and explains the allowed archive/hide action.",
            "Medium. Teachers use this feature when cleaning up unused repositories.",
            "BR-01, BR-09, BR-16, BR-17, BR-20, BR-37, BR-39",
            "Deleting a Question Bank must not remove copied question content already stored in Study Sets.",
            "The system supports archive or hide behavior when hard deletion is unsafe."
        )),
        ("2.4.5", "Add question manually", uc(
            "UC-37 Add question manually",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to manually add a question into an authorized question container so that I can build learning or assessment content for a Study Set or Question Bank. Learners can add questions only to their owned Study Sets. Teachers can add questions to owned or authorized Study Sets and Question Banks.",
            "The user clicks the Add Question action from an owned Study Set, authorized Study Set, or Teacher Question Bank.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The selected question container exists.\nPRE-4. The user has permission to author questions in the selected container.\nPRE-5. If the actor is Learner, the target container must be a Study Set owned by that Learner.\nPRE-6. If the actor is Teacher, the target container may be an owned or authorized Study Set or Question Bank.",
            "POST-1. A new question is added successfully to the selected question container.\nPOST-2. The new question appears in the question list of the selected Study Set or Question Bank.\nPOST-3. If the target container is a Study Set, the question becomes available for flashcard or quiz learning according to Study Set visibility and access rules.\nPOST-4. If the target container is a Question Bank, the question becomes reusable for future Study Set or Exam creation.\nPOST-5. If the action fails or is cancelled, no new question is saved.",
            "UC-37.0 Add Question Manually\n1. The user opens an authorized Study Set or, for Teacher only, an authorized Question Bank.\n2. The system displays the question list for the selected container.\n3. The user clicks Add Question.\n4. The system displays the Add Question form.\n5. The user enters question content, question type, answer options, correct answer, score, explanation, tags, subject, topic, chapter, lesson, and difficulty level.\n6. The user submits the question draft.\n7. The system validates the question data and the user's permission for the selected container.\n8. The system displays the question preview before saving.\n9. The user confirms saving the question.\n10. The system saves the question into the selected container.\n11. The system displays a success message.\n12. The system returns the user to the updated question list.",
            "UC-37.1 Cancel Add Question\n1. At step 5 of the Normal Flow, the user clicks Cancel.\n2. The system discards unsaved question data.\n3. The system returns the user to the question list.\n4. No new question is saved.\nUC-37.2 Save and Add Another Question\n1. At step 9, the user selects Save and Add Another.\n2. The system saves the current question.\n3. The system displays an empty Add Question form.\n4. The flow continues from step 5.\nUC-37.3 Teacher Adds Question To Question Bank\n1. At step 1, the Teacher opens an authorized Question Bank.\n2. The Teacher adds the question using the same Add Question form.\n3. The system saves the question into the Question Bank as reusable repository content.",
            "UC-37.0.E1 Invalid Question Data\n1. At step 7, the system detects missing or invalid required data.\n2. The system highlights invalid fields and displays validation messages.\n3. The user corrects the data.\n4. The flow continues from step 5.\nUC-37.0.E2 Unauthorized Question Container\n1. At step 7, the system detects that the user does not have permission for the selected container.\n2. The system blocks the action and displays MSG11.\n3. No question is saved.\nUC-37.0.E3 Container No Longer Exists\n1. At step 7 or step 10, the system detects that the selected container no longer exists or is no longer available.\n2. The system displays an error message.\n3. No question is saved.",
            "High. Learners and Teachers use this feature frequently when creating Study Sets, improving learning content, or preparing reusable questions.",
            "BR-01, BR-09, BR-18, BR-20, BR-22, BR-37, BR-38, BR-39",
            "Question authoring uses the same form and validation rules for Study Sets and Question Banks. The difference is controlled by the target container authorization. Learners cannot add questions directly to Question Banks.",
            "The system can identify the selected question container type before saving the question. Study Set questions and Question Bank questions share the same question data structure, validation rules, and preview behavior."
        )),
        ("2.4.6", "Import questions from Excel", uc(
            "UC-38 Import questions from Excel",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to import questions from an Excel template into an authorized question container so that I can create many questions efficiently.",
            "The user clicks Import Questions from an owned Study Set, authorized Study Set, or Teacher Question Bank.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The target question container exists.\nPRE-4. The user has authoring permission for the target container.\nPRE-5. The Excel file follows the platform question import template.",
            "POST-1. Valid imported questions are saved into the selected Study Set or Question Bank after user approval.\nPOST-2. Invalid rows are not saved.\nPOST-3. Row-level errors are available for review when import validation fails.\nPOST-4. If the import is cancelled, no imported question is saved.",
            "UC-38.0 Import Questions From Excel\n1. The user opens an authorized Study Set or, for Teacher only, an authorized Question Bank.\n2. The user clicks Import Questions.\n3. The system displays the import form and template requirements.\n4. The user selects or uploads an Excel file.\n5. The user confirms the target container.\n6. The system validates file format, required columns, question data, answer options, correct answers, score, and metadata.\n7. The system separates valid rows and invalid rows.\n8. The system displays import preview and row-level errors if any.\n9. The user reviews the import result.\n10. The user confirms saving valid questions.\n11. The system saves approved valid questions into the selected container.\n12. The system displays MSG28.",
            "UC-38.1 Import Valid Rows Only\n1. At step 8, the system finds both valid and invalid rows.\n2. The user chooses to save only valid rows.\n3. The system saves valid rows and keeps invalid rows unsaved.\nUC-38.2 Re-upload Corrected File\n1. At step 8, the user downloads or reviews errors.\n2. The user corrects the Excel file.\n3. The user uploads the corrected file.\n4. The flow continues from step 6.\nUC-38.3 Cancel Import\n1. At step 9, the user cancels import.\n2. The system discards the import session.\n3. No imported question is saved.",
            "UC-38.0.E1 Invalid Excel Template\n1. At step 6, the system detects that the file does not follow the required template.\n2. The system displays MSG03.\n3. No question is saved.\nUC-38.0.E2 Row-Level Validation Errors\n1. At step 6, the system detects invalid rows.\n2. The system displays MSG27 and row-level error details through UC-39.\nUC-38.0.E3 Unauthorized Target Container\n1. At step 5 or step 6, the system detects that the user cannot import into the selected container.\n2. The system blocks the import and displays MSG11.",
            "High. Users import questions when preparing large Study Sets or reusable Question Banks.",
            "BR-01, BR-09, BR-18, BR-19, BR-22, BR-37, BR-38, BR-39",
            "The same Excel template is used for Study Sets and Question Banks. Authorization is based on the selected target container.",
            "The system can validate Excel files before saving any question records."
        )),
        ("2.4.7", "View question import errors", uc(
            "UC-39 View question import errors",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to view row-level import errors so that I can correct invalid question data before saving it into a Study Set or Question Bank.",
            "The system detects invalid rows during Excel import and the user opens the import error result.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user has started UC-38 Import questions from Excel.\nPRE-3. The system has completed validation and found invalid rows.\nPRE-4. The user still has permission for the selected target container.",
            "POST-1. Import errors are displayed to the user.\nPOST-2. Invalid rows remain unsaved.\nPOST-3. The user can correct the Excel file, re-upload it, or continue with valid rows when allowed.",
            "UC-39.0 View Question Import Errors\n1. The system detects invalid rows during Excel import.\n2. The system displays MSG27.\n3. The user clicks View Errors or opens the import error panel.\n4. The system displays row number, field name, invalid value, and validation reason.\n5. The user reviews the error details.\n6. The user decides whether to correct the file, re-upload, cancel import, or save valid rows only if allowed.",
            "UC-39.1 Filter Error List\n1. At step 4, the user filters errors by row number, field name, or error type.\n2. The system updates the visible error list.\nUC-39.2 Export Error List\n1. At step 5, the user exports or copies error details if the screen supports it.\n2. The system provides the error information for correction outside the platform.",
            "UC-39.0.E1 Import Session Expired\n1. At step 3, the system detects that the import session no longer exists.\n2. The system asks the user to upload the Excel file again.\nUC-39.0.E2 Unauthorized Import Result\n1. At step 4, the system detects that the user no longer has permission for the target container.\n2. The system blocks access and displays MSG11.",
            "Medium. Users view errors when imported files contain invalid rows.",
            "BR-01, BR-09, BR-18, BR-19, BR-37",
            "This use case is always connected to UC-38 and does not save questions by itself.",
            "The system stores temporary import validation results long enough for the user to review them."
        )),
        ("2.4.8", "Preview questions before saving", uc(
            "UC-40 Preview questions before saving",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to preview manually entered, imported, copied, or AI-generated questions before saving them so that I can confirm quality and correctness.",
            "The system has question drafts from manual input, Excel import, Question Bank copy flow, or AI generation and the user opens the preview step.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The question draft exists.\nPRE-4. The target container exists.\nPRE-5. The user has permission to save questions into the target container.",
            "POST-1. Approved valid questions are saved into the selected Study Set or Question Bank.\nPOST-2. Unapproved or invalid drafts are not saved.\nPOST-3. If the target is a Study Set, saved questions are available for learning according to access rules.\nPOST-4. If the target is a Question Bank, saved questions become reusable repository content.",
            "UC-40.0 Preview Questions Before Saving\n1. The user reaches the preview step after manual entry, Excel import, Question Bank copy selection, or AI generation.\n2. The system displays each question draft with question content, options, correct answer, explanation, score, tags, subject, topic, chapter, lesson, and difficulty.\n3. The user reviews the draft questions.\n4. The user edits, removes, or confirms each draft as needed.\n5. The user clicks Save Approved Questions.\n6. The system validates approved drafts and target container authorization.\n7. The system saves approved valid questions into the selected container.\n8. The system displays a success message.",
            "UC-40.1 Edit Draft In Preview\n1. At step 4, the user edits a draft question.\n2. The system validates the edited draft.\n3. The flow continues from step 3.\nUC-40.2 Remove Draft From Preview\n1. At step 4, the user removes a draft question.\n2. The system excludes the draft from saving.\n3. The flow continues from step 3.\nUC-40.3 Save Subset Of Drafts\n1. At step 5, the user approves only selected drafts.\n2. The system saves only approved valid drafts.",
            "UC-40.0.E1 Invalid Draft Data\n1. At step 6, the system detects that an approved draft is invalid.\n2. The system highlights the invalid draft and prevents it from being saved until corrected.\nUC-40.0.E2 Unauthorized Target Container\n1. At step 6, the system detects that the user cannot save into the target container.\n2. The system blocks saving and displays MSG11.",
            "High. Preview is used whenever questions are manually created, imported, copied, or generated by AI.",
            "BR-01, BR-09, BR-18, BR-19, BR-21, BR-37, BR-39",
            "Preview is a shared quality gate before saving questions to either Study Sets or Question Banks.",
            "The system can keep temporary question drafts before final save."
        )),
        ("2.4.9", "Update question", uc(
            "UC-41 Update question",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to update an existing question in an authorized question container so that I can correct or improve its content, answer options, correct answer, score, explanation, tags, or difficulty.",
            "The user clicks Edit or Update Question from a Study Set question list, Study Set question detail, or Teacher Question Bank question list.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The account is active and not locked.\nPRE-3. The selected question container exists.\nPRE-4. The selected question exists in the container.\nPRE-5. Learners may update questions only in owned Study Sets.\nPRE-6. Teachers may update questions in owned or authorized Study Sets and Question Banks.",
            "POST-1. The selected question is updated successfully.\nPOST-2. The updated question information is displayed to the user.\nPOST-3. Existing study sessions, exam sessions, or reports that reference the question are handled according to integrity rules.\nPOST-4. If the update fails or is cancelled, the original question data remains unchanged.",
            "UC-41.0 Update Question\n1. The user opens an authorized Study Set or, for Teacher only, an authorized Question Bank.\n2. The user selects a question from the question list.\n3. The system displays the question detail page.\n4. The user clicks Edit or Update Question.\n5. The system displays the Update Question form with existing question data.\n6. The user updates question content, answer options, correct answer, score, explanation, tags, subject, topic, chapter, lesson, or difficulty level.\n7. The user submits the updated question.\n8. The system validates the updated data and authorization for the target container.\n9. The system saves the changes.\n10. The system displays a success message.\n11. The system displays the updated question detail page.",
            "UC-41.1 Cancel Update\n1. At step 6 of the Normal Flow, the user clicks Cancel.\n2. The system discards unsaved changes.\n3. The system returns the user to the question detail page.\n4. The original question data remains unchanged.\nUC-41.2 Update Metadata Only\n1. At step 6, the user updates only metadata such as tags, subject, topic, chapter, lesson, or difficulty level.\n2. The system validates that required question content and answer fields remain valid.\n3. The flow continues from step 9.\nUC-41.3 Update Correct Answer\n1. At step 6, the user changes the correct answer.\n2. The system validates that the selected correct answer belongs to the available answer options.\n3. The flow continues from step 9.",
            "UC-41.0.E1 Invalid Updated Data\n1. At step 8, the system detects invalid or missing required data.\n2. The system highlights invalid fields and displays validation messages.\n3. The user corrects the data.\n4. The flow continues from step 6.\nUC-41.0.E2 Unauthorized Question Container\n1. At step 8, the system detects that the user does not have permission to update the question in the selected container.\n2. The system blocks the action and displays MSG11.\n3. The original question remains unchanged.",
            "High. Learners and Teachers use this feature when correcting mistakes, improving question quality, or updating question metadata in authorized containers.",
            "BR-01, BR-07, BR-09, BR-18, BR-20, BR-37, BR-39",
            "If the question has already been used in exams or completed practice sessions, the system should preserve historical attempt data according to reporting and audit rules.",
            "The system can maintain question references safely when a question is updated, especially when the question is already used by Study Sets or Exam Sessions."
        )),
        ("2.4.10", "Delete question", uc(
            "UC-42 Delete question",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to delete, remove, hide, or archive a question from an authorized question container so that outdated or invalid content is no longer used.",
            "The user clicks Delete, Remove, Hide, or Archive on a question in an authorized Study Set or, for Teacher only, Question Bank.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The selected question container exists.\nPRE-4. The selected question exists in the container.\nPRE-5. The user has delete or remove permission for the selected container.",
            "POST-1. The question is removed, hidden, archived, or deleted according to integrity rules.\nPOST-2. The question is no longer available for new learning or authoring use in that container when removal is successful.\nPOST-3. Historical learning, exam, and report data remain valid.\nPOST-4. If the action fails or is cancelled, the question remains unchanged.",
            "UC-42.0 Delete Question\n1. The user opens an authorized Study Set or, for Teacher only, an authorized Question Bank.\n2. The user selects a question from the question list.\n3. The user clicks Delete, Remove, Hide, or Archive.\n4. The system displays a confirmation message.\n5. The user confirms the action.\n6. The system validates user permission and checks whether the question is referenced by practice sessions, exams, or reports.\n7. The system applies the allowed action according to integrity rules.\n8. The system displays a success message.\n9. The system displays the updated question list.",
            "UC-42.1 Cancel Deletion\n1. At step 4, the user cancels the confirmation.\n2. The system closes the confirmation message.\n3. The question remains unchanged.\nUC-42.2 Archive Or Hide Instead Of Hard Delete\n1. At step 6, the system detects historical references that must be preserved.\n2. The system archives or hides the question instead of hard-deleting it.\n3. The flow continues from step 8.\nUC-42.3 Remove Question From Study Set Only\n1. At step 7, the target container is a Study Set using copied question content.\n2. The system removes the question from the Study Set without changing the source Question Bank.",
            "UC-42.0.E1 Unauthorized Action\n1. At step 6, the system detects that the user cannot delete or remove the question from the selected container.\n2. The system blocks the action and displays MSG11.\nUC-42.0.E2 Question Locked By Integrity Rules\n1. At step 6, the system detects that the question cannot be removed because of an active exam or locked record.\n2. The system blocks hard deletion and explains the allowed archive or hide action.",
            "Medium. Users delete or remove questions when content becomes invalid, duplicated, or no longer useful.",
            "BR-01, BR-09, BR-18, BR-20, BR-37, BR-39",
            "Deleting a copied Study Set question must not grant access to or modify the source Question Bank unless the Teacher explicitly manages that Question Bank.",
            "The system supports soft delete, hide, or archive behavior when historical references exist."
        )),
        ("2.4.11", "Generate questions from material", uc(
            "UC-43 Generate questions from material",
            "Learner, Teacher",
            "Gemini API",
            "As a Premium Learner or Premium Teacher, I want to generate draft questions from learning material using AI so that I can quickly create questions for an authorized Study Set or Question Bank.",
            "The user clicks Generate Questions from Material from an owned Study Set, authorized Study Set, or Teacher Question Bank.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The user has valid Premium access.\nPRE-4. The target container exists.\nPRE-5. The user has authoring permission for the target container.\nPRE-6. The uploaded or selected learning material is valid and readable by the system.",
            "POST-1. AI-generated draft questions are displayed for review.\nPOST-2. Approved valid questions are saved into the selected Study Set or Question Bank.\nPOST-3. Unapproved or invalid generated questions are not saved.\nPOST-4. If AI generation fails or is cancelled, no generated question is saved.",
            "UC-43.0 Generate Questions From Material\n1. The user opens an authorized Study Set or, for Teacher only, an authorized Question Bank.\n2. The user clicks Generate Questions from Material.\n3. The system validates Premium access and target container authorization.\n4. The system displays the AI generation form.\n5. The user uploads or selects learning material and chooses generation settings such as question type, number of questions, subject, topic, and difficulty.\n6. The user submits the AI generation request.\n7. The system sends the request to Gemini API.\n8. Gemini API returns generated draft questions.\n9. The system displays MSG29 and opens the preview step.\n10. The user reviews, edits, removes, or approves generated drafts through UC-40.\n11. The system saves approved valid questions into the selected container.",
            "UC-43.1 Regenerate Questions\n1. At step 10, the user chooses to regenerate draft questions.\n2. The system sends a new request to Gemini API using the selected material and settings.\n3. The flow continues from step 8.\nUC-43.2 Cancel AI Generation\n1. At step 5 or step 10, the user cancels the process.\n2. The system discards unsaved generated drafts.\n3. No generated question is saved.",
            "UC-43.0.E1 Non-Premium Account\n1. At step 3, the system detects that the user does not have Premium access.\n2. The system blocks AI generation and displays MSG16.\nUC-43.0.E2 Gemini API Failure\n1. At step 7 or step 8, Gemini API fails, times out, or returns no usable content.\n2. The system displays MSG18.\n3. No generated question is saved.\nUC-43.0.E3 Unauthorized Target Container\n1. At step 3, the system detects that the user cannot save questions into the target container.\n2. The system blocks the action and displays MSG11.",
            "High for Premium users. Learners and Teachers use this feature to accelerate question creation from learning material.",
            "BR-01, BR-09, BR-18, BR-21, BR-37, BR-38, BR-39",
            "Learners may save generated questions only to owned Study Sets. Teachers may save generated questions to authorized Study Sets or Question Banks.",
            "The Gemini API can return structured draft questions that are editable before saving."
        )),
    ]

    s25 = [
        ("2.5.1", "View accessible study sets", uc(
            "UC-18 View accessible study sets",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to view Study Sets that I can access so that I can continue learning, manage my own Study Sets, or use assigned learning content.",
            "The user opens the Study Sets page or dashboard Study Sets area.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The active role is Learner or Teacher.",
            "POST-1. The system displays Study Sets accessible to the user.\nPOST-2. The user can open a Study Set detail page.\nPOST-3. Question Banks are not displayed to Learners as Study Set sources.",
            "UC-18.0 View Accessible Study Sets\n1. The user opens the Study Sets page.\n2. The system validates the user's active role and authorization.\n3. The system retrieves accessible Study Sets.\n4. For Learner, the system includes owned Study Sets, public Study Sets started by the Learner, and Study Sets assigned through joined classes.\n5. For Teacher, the system includes owned Study Sets and Study Sets authorized for managed classes.\n6. The system displays Study Set title, description, subject, topic, visibility, owner, question count, and learning status where available.\n7. The user filters, searches, sorts, or opens a Study Set detail page.",
            "UC-18.1 No Accessible Study Sets\n1. At step 3, the system finds no accessible Study Sets.\n2. The system displays an empty state.\n3. The system provides Create Study Set if the active role allows it.\nUC-18.2 Filter Study Sets\n1. At step 7, the user filters by subject, topic, ownership, assignment, or visibility.\n2. The system updates the visible Study Set list.",
            "UC-18.0.E1 Unauthorized Access\n1. At step 2, the system detects that the user is not allowed to access the requested Study Set list.\n2. The system displays MSG11.\nUC-18.0.E2 Load Failure\n1. At step 3, the system cannot load Study Set data.\n2. The system displays MSG13.",
            "High. Users access Study Sets frequently for learning and management.",
            "BR-01, BR-09, BR-22, BR-23, BR-24, BR-38, BR-39",
            "Public search for Study Sets is handled by UC-02. This use case covers authenticated accessible Study Set lists.",
            "The system can calculate Study Set accessibility based on ownership, public status, class assignment, and Teacher authorization."
        )),
        ("2.5.2", "Create study set", uc(
            "UC-44 Create study set",
            "Learner, Teacher",
            "None",
            "As a Learner or Teacher, I want to create a Study Set as the primary learning object so that I can add questions, study flashcards, take quizzes, or assign learning content to classes when authorized.",
            "The user clicks Create Study Set from the dashboard, Study Sets page, or profile learning area.",
            "PRE-1. The user has logged in successfully.\nPRE-2. The user account is active and not locked.\nPRE-3. The active role is Learner or Teacher.",
            "POST-1. A new Study Set is created successfully.\nPOST-2. The new Study Set is Public by default.\nPOST-3. The user can add questions through shared Question Management use cases.\nPOST-4. If the action fails or is cancelled, no Study Set is created.",
            "UC-44.0 Create Study Set\n1. The user opens the Create Study Set page.\n2. The system displays the Study Set form.\n3. The user enters title, description, subject, topic, and optional visibility settings.\n4. The system sets the new Study Set visibility to Public by default.\n5. The user submits the Study Set form.\n6. The system validates required Study Set information and user authorization.\n7. The system creates the Study Set.\n8. The system displays MSG35.\n9. The system opens the Study Set detail or editor page.\n10. The user adds questions through UC-37, UC-38, or UC-43 and reviews draft questions through UC-40.",
            "UC-44.1 Create Empty Study Set\n1. At step 10, the user decides not to add questions immediately.\n2. The system keeps the Study Set available as an empty Study Set until questions are added.\nUC-44.2 Teacher Copies Questions From Question Bank\n1. At step 10, the Teacher selects questions from an authorized Question Bank.\n2. The system copies selected question content into the Study Set.\n3. The copied questions become part of the Study Set.\n4. The system does not grant Learners direct access to the source Question Bank.\nUC-44.3 Change Visibility If Supported\n1. At step 3, the user changes Study Set visibility if the screen supports it.\n2. The system validates the visibility setting.\n3. The flow continues from step 5.",
            "UC-44.0.E1 Invalid Study Set Data\n1. At step 6, the system detects missing or invalid Study Set information.\n2. The system highlights invalid fields and displays validation messages.\n3. The user corrects the data.\n4. The flow continues from step 3.\nUC-44.0.E2 Unauthorized Question Bank Source\n1. During UC-44.2, the system detects that the Teacher is not authorized to use the selected Question Bank.\n2. The system blocks the copy action and displays MSG11.\nUC-44.0.E3 Non-Premium AI Generation\n1. At step 10, the user tries to generate questions with AI without Premium access.\n2. The system blocks AI generation and displays MSG16.",
            "High. Study Set creation is the main content creation flow for Learners and Teachers.",
            "BR-01, BR-09, BR-21, BR-22, BR-37, BR-38, BR-39",
            "Study Set is the primary learner-facing learning object. Question Bank is only a Teacher repository source and is not required for Study Set creation.",
            "The system supports Study Sets without questions immediately after creation, but learning modes require available questions."
        )),
        ("2.5.3", "Assign study set to class", uc(
            "UC-45 Assign study set to class",
            "Teacher",
            "Email Service",
            "As a Teacher, I want to assign an owned or authorized Study Set to a class or selected learners so that learners can study assigned content.",
            "The Teacher clicks Assign to Class from a Study Set detail page or class learning material page.",
            "PRE-1. The Teacher has logged in successfully.\nPRE-2. The Teacher account is active and not locked.\nPRE-3. The active role is Teacher.\nPRE-4. The selected Study Set exists.\nPRE-5. The Teacher owns or is authorized to assign the Study Set.\nPRE-6. The target class exists and is managed by the Teacher.",
            "POST-1. The Study Set is assigned to the selected class or learners.\nPOST-2. Eligible Learners can access the assigned Study Set.\nPOST-3. Assignment information is displayed in the class or learner Study Set list.",
            "UC-45.0 Assign Study Set To Class\n1. The Teacher opens an owned or authorized Study Set.\n2. The Teacher clicks Assign to Class.\n3. The system displays classes managed by the Teacher.\n4. The Teacher selects a class or selected learners.\n5. The Teacher confirms assignment settings.\n6. The system validates Teacher authorization for the Study Set and target class.\n7. The system creates the assignment.\n8. The system notifies eligible Learners if notification is enabled.\n9. The system displays a success message.",
            "UC-45.1 Assign To Selected Learners\n1. At step 4, the Teacher selects individual Learners instead of the full class.\n2. The system validates that selected Learners belong to the managed class.\n3. The flow continues from step 5.\nUC-45.2 Cancel Assignment\n1. At step 5, the Teacher cancels assignment.\n2. The system discards assignment changes.\n3. No Study Set assignment is created.",
            "UC-45.0.E1 Unauthorized Study Set\n1. At step 6, the system detects that the Teacher cannot assign the selected Study Set.\n2. The system blocks the action and displays MSG11.\nUC-45.0.E2 Unauthorized Class\n1. At step 6, the system detects that the Teacher does not manage the selected class.\n2. The system blocks the action and displays MSG11.",
            "High. Teachers use this feature to distribute learning materials to classes.",
            "BR-01, BR-09, BR-11, BR-15, BR-22, BR-34",
            "Only Teachers can assign Study Sets to classes. Learners can access assigned Study Sets but cannot assign them.",
            "The system can determine Teacher management scope for classes and Study Sets."
        )),
        ("2.5.4", "Study with flashcard mode", uc(
            "UC-05 Study flashcards",
            "Guest, Learner, Teacher",
            "None",
            "As a Guest, Learner, or Teacher, I want to study questions and answers using flashcard mode so that I can review learning content. Guests can study only public flashcards, while authenticated users can study public or accessible Study Sets.",
            "The user opens flashcard mode from a public Study Set or accessible Study Set detail page.",
            "PRE-1. The selected Study Set exists.\nPRE-2. The Study Set contains at least one question.\nPRE-3. Guest access is allowed only when the Study Set is Public and not hidden.\nPRE-4. Authenticated users must have access to the selected Study Set.",
            "POST-1. The user completes or exits a flashcard study session.\nPOST-2. For authenticated Learners, study progress can be saved.\nPOST-3. Guests do not receive personal progress tracking unless they register or log in.",
            "UC-05.0 Study Flashcards\n1. The user opens a public or accessible Study Set.\n2. The user clicks Study Flashcards.\n3. The system validates Study Set access.\n4. The system loads flashcards from Study Set questions.\n5. The system displays the first flashcard question.\n6. The user flips the flashcard to view the answer.\n7. The user marks the card or moves to the next flashcard.\n8. The system continues until the user completes the session or exits.\n9. The system saves progress for authenticated Learners.\n10. The system displays MSG12 when the session is completed.",
            "UC-05.1 Guest Studies Public Flashcards\n1. At step 3, the system detects that the user is Guest.\n2. The system allows flashcard mode only for Public Study Sets.\n3. The flow continues from step 4.\nUC-05.2 Shuffle Flashcards\n1. At step 4, the user selects shuffle mode if available.\n2. The system randomizes flashcard order.\n3. The flow continues from step 5.",
            "UC-05.0.E1 Unauthorized Study Set\n1. At step 3, the system detects that the user cannot access the selected Study Set.\n2. The system blocks access and displays MSG11.\nUC-05.0.E2 No Flashcards Available\n1. At step 4, the system detects that the Study Set has no available questions.\n2. The system displays an empty state and does not start flashcard mode.\nUC-05.0.E3 Load Failure\n1. At step 4, the system cannot load flashcard data.\n2. The system displays MSG13.",
            "High. Flashcard mode is a primary learning activity for public and accessible Study Sets.",
            "BR-01, BR-02, BR-09, BR-22, BR-23, BR-24",
            "Flashcard mode uses Study Set questions. Question Banks are not directly studied by Learners or Guests.",
            "The system can distinguish Guest public access from authenticated Study Set access."
        )),
        ("2.5.5", "Take study set quiz", uc(
            "UC-19 Take study set quiz",
            "Learner",
            "None",
            "As a Learner, I want to take a quiz generated from an accessible Study Set so that I can practice and assess my understanding.",
            "The Learner clicks Take Quiz from an accessible Study Set detail page.",
            "PRE-1. The Learner has logged in successfully.\nPRE-2. The Learner account is active and not locked.\nPRE-3. The selected Study Set exists and is accessible to the Learner.\nPRE-4. The Study Set contains questions available for quiz mode.",
            "POST-1. The Learner's quiz attempt is submitted.\nPOST-2. The system records score, answers, wrong answers, and progress.\nPOST-3. The Learner can review wrong answers after quiz completion.",
            "UC-19.0 Take Study Set Quiz\n1. The Learner opens an accessible Study Set.\n2. The Learner clicks Take Quiz.\n3. The system validates Learner access to the Study Set.\n4. The system generates quiz questions from the Study Set.\n5. The Learner answers each question.\n6. The Learner submits the quiz.\n7. The system grades the quiz.\n8. The system saves quiz results, wrong answers, and learning progress.\n9. The system displays MSG14 with the quiz result summary.\n10. The Learner may proceed to UC-20 Review wrong answers.",
            "UC-19.1 Exit Before Submit\n1. At step 5, the Learner exits quiz mode before submitting.\n2. The system handles unsaved answers according to configured quiz behavior.\n3. The quiz may be discarded or saved as incomplete if supported.\nUC-19.2 Retake Quiz\n1. After step 9, the Learner chooses to retake the quiz.\n2. The system generates a new quiz attempt from the accessible Study Set.\n3. The flow continues from step 5.",
            "UC-19.0.E1 Guest Attempts Quiz\n1. At step 2 or step 3, the system detects that the user is Guest.\n2. The system blocks quiz mode and displays MSG11.\nUC-19.0.E2 Unauthorized Study Set\n1. At step 3, the system detects that the Learner cannot access the selected Study Set.\n2. The system blocks quiz mode and displays MSG11.\nUC-19.0.E3 No Questions Available\n1. At step 4, the system detects that the Study Set has no questions available for quiz mode.\n2. The system displays an empty state and does not start the quiz.",
            "High. Learners use quiz mode frequently to practice Study Sets.",
            "BR-01, BR-09, BR-22, BR-23, BR-24",
            "Quiz mode is available only to authenticated Learners. Teachers may study flashcards but do not take Learner quiz attempts under this UC.",
            "The system can generate quiz questions from Study Set questions without exposing any source Question Bank."
        )),
        ("2.5.6", "Review wrong answers", uc(
            "UC-20 Review wrong answers",
            "Learner",
            "None",
            "As a Learner, I want to review questions I answered incorrectly so that I can understand mistakes and improve future performance.",
            "The Learner clicks Review Wrong Answers from quiz results, learning history, or mistake review page.",
            "PRE-1. The Learner has logged in successfully.\nPRE-2. The Learner account is active and not locked.\nPRE-3. The Learner has completed at least one quiz, practice session, or exam attempt with recorded answers.",
            "POST-1. The Learner views wrong answer details.\nPOST-2. The Learner can continue practice or request AI explanation if Premium and supported.\nPOST-3. Review activity can be considered in learning progress analytics.",
            "UC-20.0 Review Wrong Answers\n1. The Learner opens the wrong answer review page or clicks Review Wrong Answers from quiz results.\n2. The system retrieves incorrect answers from the Learner's learning history.\n3. The system displays question content, the Learner's answer, correct answer, explanation if available, Study Set or Exam source, and related metadata.\n4. The Learner reviews each wrong answer.\n5. The Learner marks the question as reviewed or moves to the next wrong answer.\n6. The system updates review status and learning progress.",
            "UC-20.1 No Wrong Answers\n1. At step 2, the system finds no wrong answers.\n2. The system displays MSG15.\nUC-20.2 Filter Wrong Answers\n1. At step 3, the Learner filters wrong answers by Study Set, Exam, subject, topic, or date.\n2. The system updates the displayed wrong answer list.\nUC-20.3 Request AI Explanation\n1. At step 4, a Premium Learner requests explanation for a question.\n2. The system continues with UC-21 Request AI answer explanation.",
            "UC-20.0.E1 Load Failure\n1. At step 2, the system cannot load wrong answer data.\n2. The system displays MSG13.\nUC-20.0.E2 Unauthorized Review Data\n1. At step 2, the system detects that the Learner is trying to access another user's review data.\n2. The system blocks access and displays MSG11.",
            "High. Learners use this feature after quizzes and exams to improve weak areas.",
            "BR-01, BR-09, BR-22, BR-24, BR-25",
            "Wrong answer review may include learner-created Study Sets, assigned Study Sets, and exam attempts when visibility rules allow review.",
            "The system records enough answer history to distinguish the Learner's answer from the correct answer."
        )),
        ("2.5.7", "Request AI answer explanation", uc(
            "UC-21 Request AI answer explanation",
            "Learner",
            "Gemini API",
            "As a Premium Learner, I want to request an AI-generated explanation for a question and answer so that I can understand the reasoning behind the correct answer.",
            "The Premium Learner clicks Request AI Explanation from a question, quiz result, or wrong answer review item.",
            "PRE-1. The Learner has logged in successfully.\nPRE-2. The Learner account is active and not locked.\nPRE-3. The Learner has valid Premium access.\nPRE-4. The selected question exists and is accessible to the Learner.\nPRE-5. Gemini API is available or the system can handle AI failure gracefully.",
            "POST-1. The system displays an AI-generated explanation if the request succeeds.\nPOST-2. The Learner can use the explanation during review.\nPOST-3. Core study, quiz, and review functions continue even if AI is unavailable.",
            "UC-21.0 Request AI Answer Explanation\n1. The Learner opens an accessible question, quiz result, or wrong answer review item.\n2. The Learner clicks Request AI Explanation.\n3. The system validates Premium access and question access.\n4. The system sends the question, answer options, correct answer, and relevant context to Gemini API.\n5. The system displays a processing message.\n6. Gemini API returns an explanation.\n7. The system displays the AI-generated explanation to the Learner.",
            "UC-21.1 Retry AI Explanation\n1. At step 7 or after an AI failure, the Learner clicks Retry if available.\n2. The system sends another request to Gemini API.\n3. The flow continues from step 5.\nUC-21.2 Continue Without AI\n1. If the Learner does not want to use AI explanation, the Learner continues reviewing manually.\n2. No AI request is sent.",
            "UC-21.0.E1 Non-Premium Account\n1. At step 3, the system detects that the Learner does not have Premium access.\n2. The system blocks the AI request and displays MSG16.\nUC-21.0.E2 Gemini API Failure\n1. At step 4 or step 6, Gemini API fails, times out, or returns no usable explanation.\n2. The system displays MSG18.\n3. The Learner may continue study or review without AI support.\nUC-21.0.E3 Unauthorized Question\n1. At step 3, the system detects that the Learner cannot access the selected question.\n2. The system displays MSG11.",
            "Medium. Premium Learners use this feature when they need additional explanation for difficult questions.",
            "BR-01, BR-09, BR-16, BR-22, BR-25",
            "This use case explains answers. It is separate from UC-43 Generate questions from material.",
            "The AI explanation is supplemental learning support and does not change the official correct answer."
        )),
        ("2.5.8", "View personal learning progress", uc(
            "UC-22 View personal learning progress",
            "Learner",
            "None",
            "As a Learner, I want to view my personal learning progress so that I can understand performance, repeated mistakes, and weak topics across Study Sets and Exams.",
            "The Learner opens the Learning Progress page or dashboard analytics area.",
            "PRE-1. The Learner has logged in successfully.\nPRE-2. The Learner account is active and not locked.\nPRE-3. The Learner has learning, quiz, practice, or exam activity, or the system can display an empty progress state.",
            "POST-1. The Learner views personal progress information.\nPOST-2. The system may show practiced questions, accuracy rate, repeated mistakes, weak topics, and study history.\nPOST-3. The Learner can choose a Study Set, topic, or wrong answer item for further review if supported.",
            "UC-22.0 View Personal Learning Progress\n1. The Learner opens the Learning Progress page.\n2. The system validates Learner authentication and authorization.\n3. The system retrieves progress data from Study Set practice, quiz attempts, wrong answer review, and exam attempts when visible.\n4. The system calculates progress metrics such as practiced questions, accuracy rate, repeated mistakes, weak topics, and study history.\n5. The system displays the learning progress dashboard.\n6. The Learner filters progress by Study Set, class, subject, topic, date range, or activity type if supported.\n7. The system updates the displayed progress data.",
            "UC-22.1 No Progress Data\n1. At step 3, the system finds no learning activity.\n2. The system displays an empty progress state.\n3. The Learner may open accessible Study Sets to start learning.\nUC-22.2 Open Weak Topic Or Mistake\n1. At step 5, the Learner selects a weak topic or repeated mistake.\n2. The system opens the related Study Set, question, or wrong answer review item when accessible.",
            "UC-22.0.E1 Load Failure\n1. At step 3 or step 4, the system cannot load or calculate progress data.\n2. The system displays MSG13.\nUC-22.0.E2 Unauthorized Progress Data\n1. At step 2, the system detects that the user is trying to view another user's progress.\n2. The system blocks access and displays MSG11.",
            "High. Learners use this feature to monitor study performance and plan review activities.",
            "BR-01, BR-09, BR-22, BR-24, BR-31",
            "Progress includes learner-created Study Sets, public Study Sets started by the Learner, assigned Study Sets, and exam results when visibility rules allow them.",
            "The system records learning activity with enough metadata to group progress by Study Set, subject, topic, and time period."
        )),
    ]
    return s24, s25


def main() -> None:
    doc = Document()
    configure(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Section 2.4 and 2.5 Use Case Specifications")
    set_font(r, size=20, bold=True, color="1F4D78")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Copy-ready content in the same field structure as the SRS UC detail template; no tables used.")
    set_font(r, size=10, color="666666")

    s24, s25 = build_data()
    doc.add_heading("2.4 Question Management & Question Bank Repository", level=1)
    for number, title_text, data in s24:
        add_uc(doc, number, title_text, data)

    doc.add_page_break()
    doc.add_heading("2.5 Study Set Creation and Learning", level=1)
    for number, title_text, data in s25:
        add_uc(doc, number, title_text, data)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
