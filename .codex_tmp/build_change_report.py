from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Bao cao thay doi SRS - StudySet Question Authoring.docx"

BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
GRAY = "F2F4F7"
BORDER = "BFBFBF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[idx])
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        set_cell_shading(cell, GRAY)
        for paragraph in cell.paragraphs:
            if paragraph.runs:
                paragraph.runs[0].bold = True
    for row_data in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_data):
            cell.text = text
    set_table_width(table, widths)
    set_table_borders(table)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, BLUE_DARK, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    footer = section.footer.paragraphs[0]
    footer.text = "Bao cao thay doi SRS - StudySet & Question Authoring"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)


def build_doc():
    doc = Document()
    configure_styles(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Bao Cao Thay Doi SRS")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("So sanh G4_SQP_SRS Document.docx voi ban Question Authoring Consolidated").italic = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Ngay lap: 01/06/2026 | Pham vi: Use Case, Business Rules, Screen Authorization, Entity, Messages, Appendix")

    doc.add_heading("1. Tong Quan", level=1)
    add_bullets(doc, [
        "So luong Use Case giu nguyen: 54 UC, tu UC-01 den UC-54.",
        "Khong con cac UC them du UC-55 den UC-62.",
        "Trong tam nghiep vu doi tu Question Bank-first sang Study Set la doi tuong hoc tap chinh.",
        "Question Bank duoc dinh nghia lai thanh kho cau hoi Teacher-only dung tai su dung.",
        "Question Authoring duoc gom thanh capability dung chung cho Study Set va Question Bank.",
        "Cac diagram/screen flow lien quan da duoc de trang trong file cuoi de nguoi dung tu sua tay.",
    ])

    doc.add_heading("2. Tom Tat Thay Doi Chinh", level=1)
    add_table(doc, ["Hang muc", "Truoc khi chinh", "Sau khi chinh"], [
        ["Study Set", "Chi Teacher tao Study Set tu selected questions.", "Learner va Teacher deu tao duoc Study Set; Study Set la learning object chinh."],
        ["Question Bank", "Noi chua cau hoi gan voi Teacher va luong tao cau hoi.", "Kho cau hoi Teacher-only dung tai su dung cho Study Set/Exam."],
        ["Question Authoring", "Gan chu yeu voi Question Bank.", "Dung chung cho authorized question container: Study Set hoac Question Bank."],
        ["AI Generation", "Premium Teacher generate question tu material.", "Premium Learner/Teacher generate question; Learner luu vao owned Study Set, Teacher luu vao Study Set/Question Bank."],
        ["Screen Auth", "Question Editor/Import/AI chi Teacher.", "Question Editor/Import/AI cho Learner va Teacher; Question Bank screens van Teacher-only."],
        ["Diagram", "Co diagram cu theo luong Question Bank-first.", "Cac anh diagram lien quan duoc de trang de sua tay."],
    ], [1.4, 2.45, 2.65])

    doc.add_heading("3. Thay Doi Theo Module", level=1)
    doc.add_heading("3.1 Record Of Changes", level=2)
    add_bullets(doc, [
        "Them dong 31/05: rework Study Set ownership va Question Bank repository use case.",
        "Them dong 01/06: gom cac Study Set question UC bi trung vao shared Question Authoring.",
    ])
    doc.add_heading("3.2 Actor", level=2)
    add_table(doc, ["Actor", "Noi dung da chinh"], [
        ["Guest", "Bo y try public question banks; doi thanh search public study sets va study public flashcards."],
        ["Learner", "Them quyen create/manage owned Study Sets, import/generate study set questions, practice quiz/flashcard."],
        ["Teacher", "Them quyen create Study Set, assign Study Set, quan ly Teacher-only Question Bank repository."],
        ["Gemini API", "Mo ta ro viec generate Study Set hoac Question Bank questions, explain answers, analyze content."],
    ], [1.4, 5.1])

    doc.add_heading("3.3 Use Case", level=2)
    add_bullets(doc, [
        "UC count giu nguyen 54; khong them UC moi trong ban cuoi.",
        "UC-02, UC-04, UC-05 mo rong actor Teacher cho public Study Set flow.",
        "UC-18 doi tu View joined study sets thanh View accessible study sets, actor Learner va Teacher.",
        "UC-33 den UC-36 giu Teacher-only nhung doi feature thanh Question Bank Repository.",
        "UC-37 den UC-43 doi thanh Question Management dung chung.",
        "UC-44 Create study set doi actor tu Teacher thanh Learner, Teacher.",
        "UC-45, UC-46, UC-53 duoc chinh mo ta de phu hop voi Study Set-first va repository logic.",
    ])
    add_table(doc, ["UC", "Thay doi"], [
        ["UC-37", "Thanh Add question manually cho Learner/Teacher trong authorized question container."],
        ["UC-38", "Thanh Import questions from Excel cho Study Set hoac Question Bank tuy quyen."],
        ["UC-39", "Thanh View question import errors dung chung."],
        ["UC-40", "Thanh Preview questions before saving dung chung cho manual/import/copied/AI."],
        ["UC-41", "Update question cho Learner/Teacher trong authorized container."],
        ["UC-42", "Delete/remove/hide/archive question trong authorized container, co integrity rules."],
        ["UC-43", "Generate questions from material cho Premium Learner/Teacher."],
        ["UC-44", "Create Study Set cho Learner/Teacher; Teacher co the copy question tu Question Bank."],
    ], [1.0, 5.5])

    doc.add_heading("3.4 Section Headings", level=2)
    add_table(doc, ["Section", "Truoc", "Sau"], [
        ["1.2.4", "Add new questions to Question Bank", "Shared Question Authoring & Question Repository"],
        ["2.4", "Question Bank Management", "Question Management & Question Bank Repository"],
        ["2.5", "Study Set Learning", "Study Set Creation and Learning"],
        ["3.4", "Question Bank Management", "Question Management & Question Bank Repository"],
        ["3.5", "Study Set Learning", "Study Set Creation and Learning"],
    ], [1.0, 2.65, 2.85])

    doc.add_heading("3.5 Screen Authorization", level=2)
    add_bullets(doc, [
        "Question Editor, Import Questions, AI Question Generation doi tu Teacher-only thanh Learner va Teacher.",
        "Question Banks, Question Bank Detail, Question Bank Edit van Teacher-only.",
        "Bo cac screen du: Study Set Editor, Study Set Excel Import, Study Set AI Generation, Create Study Set from Question Bank.",
        "Learner Study Sets va Teacher Study Sets duoc gom thanh Accessible Study Sets va Study Set Management.",
    ])

    doc.add_heading("3.6 Entities", level=2)
    add_table(doc, ["Entity", "Noi dung chinh da doi"], [
        ["Learner", "Co the create/manage owned Study Sets."],
        ["Teacher", "Quan ly Teacher-only Question Bank, Study Set, Exam, Report."],
        ["Question Bank", "Teacher-owned reusable repository; Learner/Guest khong truy cap truc tiep."],
        ["Question", "Item nam trong authorized container: Study Set hoac Question Bank."],
        ["Study Set", "Primary learner-facing object, Public mac dinh, chua questions tu shared Question Management."],
        ["AI Interaction", "Ho tro generate Study Set/Question Bank questions."],
    ], [1.4, 5.1])

    doc.add_heading("3.7 Business Rules", level=2)
    add_bullets(doc, [
        "BR-16 den BR-20 doi tu Question Bank-only sang Question Management/repository logic.",
        "BR-21 mo AI generation cho Premium Learner va Premium Teacher.",
        "BR-22 mo ta quyen access Study Set: Public, owned, assigned, authorized.",
        "BR-37 them rule shared Question Authoring.",
        "BR-38 them rule Study Set moi Public mac dinh.",
        "BR-39 them rule copy question tu restricted Question Bank khong cap quyen Question Bank cho Learner.",
    ])

    doc.add_heading("3.8 System Messages Va Appendix", level=2)
    add_bullets(doc, [
        "MSG02, MSG11, MSG13, MSG16, MSG27, MSG28, MSG29, MSG38 duoc chinh de phu hop Study Set va Question Management chung.",
        "Them MSG35: Study Set created successfully and published with Public visibility.",
        "Them MSG42: Question banks are available to Teachers only.",
        "SEC-04 cap nhat Premium AI cho Study Sets hoac Question Banks.",
        "DI-01 them validation cho Study Set format va question container authorization.",
        "DI-04 them rang buoc khong lam lo restricted Question Bank qua Study Set access.",
    ])

    doc.add_heading("4. Ket Qua Kiem Tra", level=1)
    add_table(doc, ["Hang muc", "Ket qua"], [
        ["UC count", "54 UC, tu UC-01 den UC-54."],
        ["UC du", "Khong con UC-55 den UC-62."],
        ["Question Management", "UC-37 den UC-43 actor Learner, Teacher va feature Question Management."],
        ["Question Bank", "Question Bank screens va metadata UC van Teacher-only."],
        ["Study Set", "Learner va Teacher deu co Create Study Set."],
        ["Diagram", "Anh diagram lien quan de trang de sua tay."],
        ["Visual render", "Chua render duoc vi moi truong thieu LibreOffice/soffice."],
    ], [2.0, 4.5])

    doc.add_heading("5. Ket Luan", level=1)
    doc.add_paragraph(
        "Ban chinh sua da dua Study Set ve dung vai tro la doi tuong hoc tap chinh, "
        "dong thoi giu Question Bank nhu kho cau hoi dung lai cho Teacher. "
        "Cac thao tac tao/sua/xoa/import/AI generate question duoc gom thanh Question Management chung, "
        "giup SRS gon hon va tranh lap use case giua Study Set va Question Bank."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
