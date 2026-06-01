# -*- coding: utf-8 -*-
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "G4_SQP_SRS Document - Question Authoring Consolidated Blank Diagrams.docx"
OUT = ROOT / "docs" / "G4_SQP_SRS Document - Switch Role Detailed Updated.docx"


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Calibri"
            if run._element.rPr is not None:
                run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
                run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def clone_row(table):
    new_tr = deepcopy(table.rows[-1]._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def append_table_row(table, values: list[str]) -> None:
    row = clone_row(table)
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def fill_first_empty_row(table, values: list[str]) -> bool:
    for row in table.rows[1:]:
        if all(not cell.text.strip() for cell in row.cells):
            for cell, value in zip(row.cells, values):
                set_cell_text(cell, value)
            return True
    return False


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def add_detail_section_after(doc: Document, anchor_text: str) -> None:
    anchor = next(p for p in doc.paragraphs if p.text.strip() == anchor_text)
    current = anchor
    lines = [
        ("2.2.9 Switch active role", "Heading 4"),
        ("Actor: Learner, Teacher", None),
        ("Preconditions:", None),
        ("1. User has logged in successfully.", None),
        ("2. User account is active.", None),
        ("3. User is currently using the Learner or Teacher role.", None),
        ("Main Flow:", None),
        ("1. User opens the Edit personal profile page.", None),
        ("2. System displays the current active role.", None),
        ("3. User selects the other available role: Learner or Teacher.", None),
        ("4. System updates the active role for the session and account context.", None),
        ("5. System displays MSG43.", None),
        ("6. System redirects the user to the dashboard that matches the new active role.", None),
        ("Postconditions:", None),
        ("1. The active role is updated.", None),
        ("2. UI and API permissions are applied according to the new active role.", None),
        ("3. Personal data, Study Sets, Classes, Question Banks, and Exams are not deleted or transferred.", None),
        ("Notes:", None),
        ("1. Switch Role is a control inside the Edit personal profile page, not a separate screen.", None),
        ("2. Users cannot switch to Admin through this use case.", None),
    ]
    for text, style in lines:
        current = insert_paragraph_after(current, text, style)


def add_lines_after_heading(doc: Document, heading_text: str, lines: list[str]) -> None:
    anchor = next(p for p in doc.paragraphs if p.text.strip() == heading_text)
    current = anchor
    for line in lines:
        current = insert_paragraph_after(current, line, None)


def normalize_empty_heading_paragraphs(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if not paragraph.text.strip() and paragraph.style.name.startswith("Heading"):
            paragraph.style = doc.styles["normal"]


def add_question_management_details(doc: Document) -> None:
    details = {
        "2.4.1 View question banks": [
            "Detailed Specification: Teacher views Question Banks as reusable question repositories, not as the primary learning flow.",
            "Actor: Teacher.",
            "Preconditions: Teacher is logged in and active; Teacher has created or is authorized to access at least one Question Bank.",
            "Main Flow: Teacher opens Question Banks, system lists owned or authorized repositories with title, subject, topic, visibility, question count, and last updated time.",
            "Postconditions: Teacher can choose a Question Bank to manage metadata, author questions, or reuse questions for Study Set or Exam creation.",
            "Authorization Note: Learner and Guest cannot directly view Question Banks.",
        ],
        "2.4.2 Create question bank": [
            "Detailed Specification: Teacher creates only the reusable repository metadata here; question content is added later through shared Question Management use cases.",
            "Actor: Teacher.",
            "Preconditions: Teacher is logged in and active.",
            "Main Flow: Teacher enters repository title, description, subject, topic, visibility, and configuration; system validates required fields; system creates the Question Bank.",
            "Postconditions: Empty Question Bank repository exists and can receive questions through UC-37, UC-38, or UC-43.",
            "Authorization Note: Learner cannot create Question Banks; Learner creates Study Sets instead.",
        ],
        "2.4.3 Update question bank information": [
            "Detailed Specification: Teacher updates repository metadata only; question content is managed by shared Question Management use cases.",
            "Actor: Teacher.",
            "Preconditions: Teacher owns or is authorized to update the selected Question Bank.",
            "Main Flow: Teacher edits title, description, subject, topic, visibility, or repository configuration; system validates and saves the metadata.",
            "Postconditions: Question Bank metadata is updated without changing existing question content unless Teacher separately uses UC-37 to UC-43.",
            "Authorization Note: Metadata update does not grant Learner direct access to the Question Bank.",
        ],
        "2.4.4 Delete question bank": [
            "Detailed Specification: Teacher deletes, archives, or hides a Question Bank repository subject to integrity rules.",
            "Actor: Teacher.",
            "Preconditions: Teacher owns or is authorized to delete/archive the selected Question Bank.",
            "Main Flow: Teacher selects delete/archive; system checks whether questions are referenced by exams, reports, or copied Study Sets; system applies the allowed action.",
            "Postconditions: Repository is deleted, archived, or hidden; historical exam, learning, and report records remain valid.",
            "Authorization Note: Learner cannot delete or archive Question Banks.",
        ],
        "2.4.5 Add question manually": [
            "Detailed Specification: Learner or Teacher manually creates a question inside an authorized question container.",
            "Actors: Learner, Teacher.",
            "Target Containers: Learner may target only owned Study Sets. Teacher may target owned or authorized Study Sets and Question Banks.",
            "Preconditions: User is logged in; account is active; target Study Set or Question Bank exists; user has authoring permission for the target container.",
            "Main Flow: User selects Add Question, enters question content, type, answer options, correct answer, score, explanation, tags, subject/topic/chapter/lesson, and difficulty; system validates data; user previews; system saves the question into the selected container.",
            "Postconditions: Question is stored in the selected Study Set or Question Bank only; saving into a Study Set does not expose any Teacher Question Bank.",
            "Exceptions: Missing required fields use MSG02/MSG03; unauthorized target uses MSG11.",
        ],
        "2.4.6 Import questions from Excel": [
            "Detailed Specification: Learner or Teacher imports questions from the shared Excel template into an authorized question container.",
            "Actors: Learner, Teacher.",
            "Target Containers: Learner may import only into owned Study Sets. Teacher may import into authorized Study Sets or Question Banks.",
            "Preconditions: User is logged in; target container exists; Excel file follows the required template.",
            "Main Flow: User selects target container, uploads Excel file, system validates file structure and row data, system separates valid and invalid rows, user reviews preview/error results, and valid questions are saved after approval.",
            "Postconditions: Valid questions are stored in the selected Study Set or Question Bank; invalid rows are not saved.",
            "Exceptions: Invalid template uses MSG03; row-level errors use MSG27; successful save uses MSG28.",
        ],
        "2.4.7 View question import errors": [
            "Detailed Specification: User reviews row-level errors from an Excel import before saving valid questions.",
            "Actors: Learner, Teacher.",
            "Preconditions: User has uploaded an Excel file and the system has found invalid rows.",
            "Main Flow: System displays row number, field name, invalid value, and validation reason; user corrects the file or proceeds with valid rows when allowed.",
            "Postconditions: Invalid rows remain unsaved; user understands which rows must be fixed.",
            "Authorization Note: Error display follows the same target-container permission as UC-38.",
        ],
        "2.4.8 Preview questions before saving": [
            "Detailed Specification: User previews manually entered, imported, copied, or AI-generated questions before saving them to a Study Set or Question Bank.",
            "Actors: Learner, Teacher.",
            "Target Containers: Learner previews questions only for owned Study Sets. Teacher previews questions for authorized Study Sets or Question Banks.",
            "Preconditions: Question draft exists from manual input, Excel import, Question Bank copy flow, or AI generation.",
            "Main Flow: System displays question content, options, correct answer, explanation, score, tags, and metadata; user edits/removes draft items; user confirms save; system validates and persists approved questions.",
            "Postconditions: Only approved valid questions are saved into the selected container.",
            "Exceptions: Invalid question draft remains in preview until corrected or removed.",
        ],
        "2.4.9 Update question": [
            "Detailed Specification: Learner or Teacher updates an existing question inside an authorized question container.",
            "Actors: Learner, Teacher.",
            "Target Containers: Learner may update questions only in owned Study Sets. Teacher may update questions in owned or authorized Study Sets and Question Banks.",
            "Preconditions: Target container and question exist; user has update permission for that container.",
            "Main Flow: User opens the question, selects Edit, changes content/options/correct answer/explanation/score/tags/metadata, system validates, and system saves the updated question.",
            "Postconditions: Updated question appears in the selected container; historical learning, exam, and report integrity is preserved.",
            "Exceptions: Unauthorized update uses MSG11; invalid data uses MSG02/MSG03.",
        ],
        "2.4.10 Delete question": [
            "Detailed Specification: Learner or Teacher removes a question from an authorized question container subject to integrity rules.",
            "Actors: Learner, Teacher.",
            "Target Containers: Learner may delete/remove questions only from owned Study Sets. Teacher may delete/remove/hide/archive questions from authorized Study Sets and Question Banks.",
            "Preconditions: Target container and question exist; user has delete/remove permission.",
            "Main Flow: User selects Delete/Remove, system checks references in practice sessions, exams, and reports, system asks for confirmation when needed, and system deletes, hides, archives, or removes the question according to integrity rules.",
            "Postconditions: Question is no longer available for new learning/authoring use in that container; historical records remain valid.",
            "Exceptions: If the question is locked by an official exam or historical record, system archives/hides instead of hard-deleting.",
        ],
        "2.4.11 Generate questions from material": [
            "Detailed Specification: Premium Learner or Premium Teacher uses AI to generate draft questions from uploaded learning material, then saves approved questions into an authorized question container.",
            "Actors: Learner, Teacher.",
            "Target Containers: Learner may save generated questions only to owned Study Sets. Teacher may save generated questions to authorized Study Sets or Question Banks.",
            "Preconditions: User is Premium; material file/content is valid; target container exists; user has authoring permission.",
            "Main Flow: User selects target container, uploads or selects learning material, chooses generation settings, system sends request to Gemini API, system returns generated draft questions, user previews/edits/removes drafts, and approved questions are saved.",
            "Postconditions: AI-generated questions are saved only after user approval through preview.",
            "Exceptions: Non-Premium access uses MSG16; AI timeout/failure uses MSG18; invalid generated drafts must be corrected before saving.",
        ],
        "2.5.2 Create study set": [
            "Detailed Specification: Learner or Teacher creates a Study Set as the primary learner-facing learning object, then adds questions through shared Question Management.",
            "Actors: Learner, Teacher.",
            "Preconditions: User is logged in and active.",
            "Main Flow: User enters Study Set title, description, subject, topic, and visibility; system sets new Study Set visibility to Public by default; system creates the Study Set; user adds questions through UC-37, UC-38, or UC-43 and reviews drafts through UC-40.",
            "Teacher Optional Flow: Teacher may copy selected questions from an authorized Question Bank into the Study Set; copied question content becomes part of the Study Set and does not grant Learners direct access to the source Question Bank.",
            "Postconditions: Study Set exists as a public learning resource by default and can be studied, quizzed, assigned by Teacher, or discovered through public search depending on authorization rules.",
            "Authorization Note: Learner cannot create Study Sets from Teacher Question Banks directly.",
        ],
    }
    for heading, lines in details.items():
        add_lines_after_heading(doc, heading, lines)


def update_uc41_detail_table(doc: Document) -> None:
    table = doc.tables[6]
    replacements = {
        "Teacher clicks": "user clicks",
        "Teacher to the question detail page": "user to the question detail page",
        "Teacher updates": "user updates",
        "Teacher changes": "user changes",
        "The Teacher corrects": "The user corrects",
    }
    for row in table.rows:
        for cell in row.cells:
            text = cell.text
            new_text = text
            for old, new in replacements.items():
                new_text = new_text.replace(old, new)
            if new_text != text:
                set_cell_text(cell, new_text)


def main() -> None:
    doc = Document(SRC)

    # Context summary.
    replace_paragraph_text(
        doc.paragraphs[29],
        "The Smart Quiz Platform (SQP) is an integrated online learning and assessment system designed to support the full cycle of study, practice, and examination. Guests can discover and study public study sets, then register to become a Learner or Teacher. Learners can join classes via code or invitation link, create their own public study sets, add questions manually, import questions from Excel, generate questions from learning materials with Premium AI access, study through flashcard and quiz modes, review wrong answers, and track personal learning progress. Learners and Teachers can switch their active role from the User Profile page. Teachers have the same study set creation capabilities as Learners, and can also assign study sets to classes, create study sets from their reusable question banks, configure official exam sessions, and export scoreboards and analytics reports. Question Banks now act as Teacher-only reusable repositories for questions that can be used repeatedly when creating study sets or exam sessions; they are not the main learner-facing learning object. Both Learners and Teachers can upgrade to Premium through an integrated Payment Gateway. Admins oversee user accounts, role permissions, public learning resources, and system stability. Email Service handles automated notifications, and Supabase Authentication handles login, session management, and password resets.",
    )

    # Record of changes.
    record_values = [
        "01/06",
        "M, A",
        "1.2, 1.3, 2.2, 5.1",
        "Added UC-55 Switch active role for Learner/Teacher profile role switching",
    ]
    if not fill_first_empty_row(doc.tables[0], record_values):
        append_table_row(doc.tables[0], record_values)

    # Actor descriptions.
    actor_table = doc.tables[1]
    set_cell_text(
        actor_table.rows[2].cells[2],
        "A registered learner who joins classes, creates and manages owned study sets, imports or generates study set questions, practices flashcards and quizzes, participates in exams, reviews mistakes, tracks learning progress, and can switch the active role to Teacher from User Profile.",
    )
    set_cell_text(
        actor_table.rows[3].cells[2],
        "An educator who creates and manages classes, creates study sets, assigns study sets to classes, manages Teacher-only question banks as reusable repositories, configures exam sessions, exports scoreboards or analytics reports, and can switch the active role to Learner from User Profile.",
    )

    # UC list.
    append_table_row(
        doc.tables[2],
        [
            "UC-55",
            "Switch active role",
            "Learner, Teacher",
            "User Profile / Role Switching",
            "User switches the active role between Learner and Teacher in the same account.",
        ],
    )

    # UC detail heading and content.
    add_detail_section_after(doc, "2.2.8 Change password")

    # Detailed use case specifications for shared Question Management and Study Set creation.
    add_question_management_details(doc)
    update_uc41_detail_table(doc)
    normalize_empty_heading_paragraphs(doc)

    # Entity model.
    entity_table = doc.tables[5]
    set_cell_text(
        entity_table.rows[1].cells[2],
        "A student or learner using Learner as the active role, who joins classes, creates and manages owned study sets, practices study sets, takes exams, submits answers, and uses learning features in the platform.",
    )
    set_cell_text(
        entity_table.rows[2].cells[2],
        "An educator using Teacher as the active role, who creates and manages classes, study sets, Teacher-only question banks, exams, and reports.",
    )
    append_table_row(
        entity_table,
        [
            "17",
            "User Account",
            "A registered account with availableRoles for Learner and Teacher self role switching, an activeRole that controls the current session permission context, and protected Admin role permissions that cannot be self-assigned.",
        ],
    )

    # Business rules.
    br_table = doc.tables[10]
    set_cell_text(
        br_table.rows[9].cells[2],
        "Each authenticated user must have a valid active role. The active role determines the user's accessible features and permissions for the current session. Learner and Teacher users may switch active role between Learner and Teacher through the User Profile page, while Admin role assignment remains restricted to Admin management.",
    )
    set_cell_text(
        br_table.rows[10].cells[2],
        "Users may update only allowed profile information such as full name, phone number, avatar, or profile description. They must not directly modify protected fields such as Admin role, account status, or system permissions; switching active role between Learner and Teacher is allowed only through the Switch active role use case.",
    )
    append_table_row(
        br_table,
        [
            "BR-40",
            "Role Switching",
            "A regular account may switch active role between Learner and Teacher. When activeRole is Learner, only Learner permissions apply. When activeRole is Teacher, only Teacher permissions apply. Users cannot self-switch to Admin, and switching role must not transfer ownership or delete existing data.",
        ],
    )

    # Messages.
    append_table_row(
        doc.tables[11],
        [
            "MSG43",
            "Success",
            "User switches active role successfully",
            "Role switched successfully.",
        ],
    )

    # Security consistency.
    replace_paragraph_text(
        doc.paragraphs[279],
        'Role-Based Access Control (RBAC): The system must strictly enforce active role privileges. Users with activeRole "Learner" must be completely blocked from backend API endpoints or UI views reserved for "Teacher" or "Admin" roles, and users with activeRole "Teacher" must be blocked from "Admin" features, returning an HTTP 403 Forbidden status code.',
    )
    sec_table = doc.tables[12]
    set_cell_text(
        sec_table.rows[2].cells[2],
        "The system must enforce Role-Based Access Control. Access to features must be restricted based on the current active role: Learner, Teacher, or Admin.",
    )

    # Save a new version instead of overwriting the previous SRS deliverable.
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
