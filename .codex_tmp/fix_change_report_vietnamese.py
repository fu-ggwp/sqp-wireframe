# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "Bao cao thay doi SRS - StudySet Question Authoring.docx"

TEXT_MAP = {
    "Bao Cao Thay Doi SRS": "Báo cáo thay đổi SRS",
    "So sanh G4_SQP_SRS Document.docx voi ban Question Authoring Consolidated": "So sánh G4_SQP_SRS Document.docx với bản Question Authoring Consolidated",
    "Ngay lap: 01/06/2026 | Pham vi: Use Case, Business Rules, Screen Authorization, Entity, Messages, Appendix": "Ngày lập: 01/06/2026 | Phạm vi: Use Case, Business Rules, Screen Authorization, Entity, Messages, Appendix",
    "1. Tong Quan": "1. Tổng quan",
    "So luong Use Case giu nguyen: 54 UC, tu UC-01 den UC-54.": "Số lượng Use Case giữ nguyên: 54 UC, từ UC-01 đến UC-54.",
    "Khong con cac UC them du UC-55 den UC-62.": "Không còn các UC thêm dư UC-55 đến UC-62.",
    "Trong tam nghiep vu doi tu Question Bank-first sang Study Set la doi tuong hoc tap chinh.": "Trọng tâm nghiệp vụ đổi từ Question Bank-first sang Study Set là đối tượng học tập chính.",
    "Question Bank duoc dinh nghia lai thanh kho cau hoi Teacher-only dung tai su dung.": "Question Bank được định nghĩa lại thành kho câu hỏi Teacher-only dùng tái sử dụng.",
    "Question Authoring duoc gom thanh capability dung chung cho Study Set va Question Bank.": "Question Authoring được gom thành capability dùng chung cho Study Set và Question Bank.",
    "Cac diagram/screen flow lien quan da duoc de trang trong file cuoi de nguoi dung tu sua tay.": "Các diagram/screen flow liên quan đã được để trắng trong file cuối để người dùng tự sửa tay.",
    "2. Tom Tat Thay Doi Chinh": "2. Tóm tắt thay đổi chính",
    "3. Thay Doi Theo Module": "3. Thay đổi theo module",
    "3.1 Record Of Changes": "3.1 Record Of Changes",
    "Them dong 31/05: rework Study Set ownership va Question Bank repository use case.": "Thêm dòng 31/05: rework Study Set ownership và Question Bank repository use case.",
    "Them dong 01/06: gom cac Study Set question UC bi trung vao shared Question Authoring.": "Thêm dòng 01/06: gom các Study Set question UC bị trùng vào shared Question Authoring.",
    "UC count giu nguyen 54; khong them UC moi trong ban cuoi.": "UC count giữ nguyên 54; không thêm UC mới trong bản cuối.",
    "UC-02, UC-04, UC-05 mo rong actor Teacher cho public Study Set flow.": "UC-02, UC-04, UC-05 mở rộng actor Teacher cho public Study Set flow.",
    "UC-18 doi tu View joined study sets thanh View accessible study sets, actor Learner va Teacher.": "UC-18 đổi từ View joined study sets thành View accessible study sets, actor Learner và Teacher.",
    "UC-33 den UC-36 giu Teacher-only nhung doi feature thanh Question Bank Repository.": "UC-33 đến UC-36 giữ Teacher-only nhưng đổi feature thành Question Bank Repository.",
    "UC-37 den UC-43 doi thanh Question Management dung chung.": "UC-37 đến UC-43 đổi thành Question Management dùng chung.",
    "UC-44 Create study set doi actor tu Teacher thanh Learner, Teacher.": "UC-44 Create study set đổi actor từ Teacher thành Learner, Teacher.",
    "UC-45, UC-46, UC-53 duoc chinh mo ta de phu hop voi Study Set-first va repository logic.": "UC-45, UC-46, UC-53 được chỉnh mô tả để phù hợp với Study Set-first và repository logic.",
    "3.5 Screen Authorization": "3.5 Screen Authorization",
    "Question Editor, Import Questions, AI Question Generation doi tu Teacher-only thanh Learner va Teacher.": "Question Editor, Import Questions, AI Question Generation đổi từ Teacher-only thành Learner và Teacher.",
    "Question Banks, Question Bank Detail, Question Bank Edit van Teacher-only.": "Question Banks, Question Bank Detail, Question Bank Edit vẫn Teacher-only.",
    "Bo cac screen du: Study Set Editor, Study Set Excel Import, Study Set AI Generation, Create Study Set from Question Bank.": "Bỏ các screen dư: Study Set Editor, Study Set Excel Import, Study Set AI Generation, Create Study Set from Question Bank.",
    "Learner Study Sets va Teacher Study Sets duoc gom thanh Accessible Study Sets va Study Set Management.": "Learner Study Sets và Teacher Study Sets được gom thành Accessible Study Sets và Study Set Management.",
    "3.7 Business Rules": "3.7 Business Rules",
    "BR-16 den BR-20 doi tu Question Bank-only sang Question Management/repository logic.": "BR-16 đến BR-20 đổi từ Question Bank-only sang Question Management/repository logic.",
    "BR-21 mo AI generation cho Premium Learner va Premium Teacher.": "BR-21 mở AI generation cho Premium Learner và Premium Teacher.",
    "BR-22 mo ta quyen access Study Set: Public, owned, assigned, authorized.": "BR-22 mô tả quyền access Study Set: Public, owned, assigned, authorized.",
    "BR-37 them rule shared Question Authoring.": "BR-37 thêm rule shared Question Authoring.",
    "BR-38 them rule Study Set moi Public mac dinh.": "BR-38 thêm rule Study Set mới Public mặc định.",
    "BR-39 them rule copy question tu restricted Question Bank khong cap quyen Question Bank cho Learner.": "BR-39 thêm rule copy question từ restricted Question Bank không cấp quyền Question Bank cho Learner.",
    "3.8 System Messages Va Appendix": "3.8 System Messages và Appendix",
    "MSG02, MSG11, MSG13, MSG16, MSG27, MSG28, MSG29, MSG38 duoc chinh de phu hop Study Set va Question Management chung.": "MSG02, MSG11, MSG13, MSG16, MSG27, MSG28, MSG29, MSG38 được chỉnh để phù hợp Study Set và Question Management chung.",
    "Them MSG35: Study Set created successfully and published with Public visibility.": "Thêm MSG35: Study Set created successfully and published with Public visibility.",
    "Them MSG42: Question banks are available to Teachers only.": "Thêm MSG42: Question banks are available to Teachers only.",
    "SEC-04 cap nhat Premium AI cho Study Sets hoac Question Banks.": "SEC-04 cập nhật Premium AI cho Study Sets hoặc Question Banks.",
    "DI-01 them validation cho Study Set format va question container authorization.": "DI-01 thêm validation cho Study Set format và question container authorization.",
    "DI-04 them rang buoc khong lam lo restricted Question Bank qua Study Set access.": "DI-04 thêm ràng buộc không làm lộ restricted Question Bank qua Study Set access.",
    "4. Ket Qua Kiem Tra": "4. Kết quả kiểm tra",
    "5. Ket Luan": "5. Kết luận",
    "Ban chinh sua da dua Study Set ve dung vai tro la doi tuong hoc tap chinh, dong thoi giu Question Bank nhu kho cau hoi dung lai cho Teacher. Cac thao tac tao/sua/xoa/import/AI generate question duoc gom thanh Question Management chung, giup SRS gon hon va tranh lap use case giua Study Set va Question Bank.": "Bản chỉnh sửa đã đưa Study Set về đúng vai trò là đối tượng học tập chính, đồng thời giữ Question Bank như kho câu hỏi dùng lại cho Teacher. Các thao tác tạo/sửa/xóa/import/AI generate question được gom thành Question Management chung, giúp SRS gọn hơn và tránh lặp use case giữa Study Set và Question Bank.",
    "Hang muc": "Hạng mục",
    "Truoc khi chinh": "Trước khi chỉnh",
    "Sau khi chinh": "Sau khi chỉnh",
    "Chi Teacher tao Study Set tu selected questions.": "Chỉ Teacher tạo Study Set từ selected questions.",
    "Learner va Teacher deu tao duoc Study Set; Study Set la learning object chinh.": "Learner và Teacher đều tạo được Study Set; Study Set là learning object chính.",
    "Noi chua cau hoi gan voi Teacher va luong tao cau hoi.": "Nơi chứa câu hỏi gắn với Teacher và luồng tạo câu hỏi.",
    "Kho cau hoi Teacher-only dung tai su dung cho Study Set/Exam.": "Kho câu hỏi Teacher-only dùng tái sử dụng cho Study Set/Exam.",
    "Gan chu yeu voi Question Bank.": "Gắn chủ yếu với Question Bank.",
    "Dung chung cho authorized question container: Study Set hoac Question Bank.": "Dùng chung cho authorized question container: Study Set hoặc Question Bank.",
    "Premium Teacher generate question tu material.": "Premium Teacher generate question từ material.",
    "Premium Learner/Teacher generate question; Learner luu vao owned Study Set, Teacher luu vao Study Set/Question Bank.": "Premium Learner/Teacher generate question; Learner lưu vào owned Study Set, Teacher lưu vào Study Set/Question Bank.",
    "Question Editor/Import/AI chi Teacher.": "Question Editor/Import/AI chỉ Teacher.",
    "Question Editor/Import/AI cho Learner va Teacher; Question Bank screens van Teacher-only.": "Question Editor/Import/AI cho Learner và Teacher; Question Bank screens vẫn Teacher-only.",
    "Co diagram cu theo luong Question Bank-first.": "Có diagram cũ theo luồng Question Bank-first.",
    "Cac anh diagram lien quan duoc de trang de sua tay.": "Các ảnh diagram liên quan được để trắng để sửa tay.",
    "Noi dung da chinh": "Nội dung đã chỉnh",
    "Bo y try public question banks; doi thanh search public study sets va study public flashcards.": "Bỏ ý 'try public question banks'; đổi thành search public study sets và study public flashcards.",
    "Them quyen create/manage owned Study Sets, import/generate study set questions, practice quiz/flashcard.": "Thêm quyền create/manage owned Study Sets, import/generate study set questions, practice quiz/flashcard.",
    "Them quyen create Study Set, assign Study Set, quan ly Teacher-only Question Bank repository.": "Thêm quyền create Study Set, assign Study Set, quản lý Teacher-only Question Bank repository.",
    "Mo ta ro viec generate Study Set hoac Question Bank questions, explain answers, analyze content.": "Mô tả rõ việc generate Study Set hoặc Question Bank questions, explain answers, analyze content.",
    "Thay doi": "Thay đổi",
    "Thanh Add question manually cho Learner/Teacher trong authorized question container.": "Thành Add question manually cho Learner/Teacher trong authorized question container.",
    "Thanh Import questions from Excel cho Study Set hoac Question Bank tuy quyen.": "Thành Import questions from Excel cho Study Set hoặc Question Bank tùy quyền.",
    "Thanh View question import errors dung chung.": "Thành View question import errors dùng chung.",
    "Thanh Preview questions before saving dung chung cho manual/import/copied/AI.": "Thành Preview questions before saving dùng chung cho manual/import/copied/AI.",
    "Update question cho Learner/Teacher trong authorized container.": "Update question cho Learner/Teacher trong authorized container.",
    "Delete/remove/hide/archive question trong authorized container, co integrity rules.": "Delete/remove/hide/archive question trong authorized container, có integrity rules.",
    "Generate questions from material cho Premium Learner/Teacher.": "Generate questions from material cho Premium Learner/Teacher.",
    "Create Study Set cho Learner/Teacher; Teacher co the copy question tu Question Bank.": "Create Study Set cho Learner/Teacher; Teacher có thể copy question từ Question Bank.",
    "Truoc": "Trước",
    "Noi dung chinh da doi": "Nội dung chính đã đổi",
    "Co the create/manage owned Study Sets.": "Có thể create/manage owned Study Sets.",
    "Quan ly Teacher-only Question Bank, Study Set, Exam, Report.": "Quản lý Teacher-only Question Bank, Study Set, Exam, Report.",
    "Teacher-owned reusable repository; Learner/Guest khong truy cap truc tiep.": "Teacher-owned reusable repository; Learner/Guest không truy cập trực tiếp.",
    "Item nam trong authorized container: Study Set hoac Question Bank.": "Item nằm trong authorized container: Study Set hoặc Question Bank.",
    "Primary learner-facing object, Public mac dinh, chua questions tu shared Question Management.": "Primary learner-facing object, Public mặc định, chứa questions từ shared Question Management.",
    "Ho tro generate Study Set/Question Bank questions.": "Hỗ trợ generate Study Set/Question Bank questions.",
    "Ket qua": "Kết quả",
    "54 UC, tu UC-01 den UC-54.": "54 UC, từ UC-01 đến UC-54.",
    "UC du": "UC dư",
    "Khong con UC-55 den UC-62.": "Không còn UC-55 đến UC-62.",
    "UC-37 den UC-43 actor Learner, Teacher va feature Question Management.": "UC-37 đến UC-43 actor Learner, Teacher và feature Question Management.",
    "Question Bank screens va metadata UC van Teacher-only.": "Question Bank screens và metadata UC vẫn Teacher-only.",
    "Learner va Teacher deu co Create Study Set.": "Learner và Teacher đều có Create Study Set.",
    "Anh diagram lien quan de trang de sua tay.": "Ảnh diagram liên quan để trắng để sửa tay.",
    "Chua render duoc vi moi truong thieu LibreOffice/soffice.": "Chưa render được vì môi trường thiếu LibreOffice/soffice.",
    "Bao cao thay doi SRS - StudySet & Question Authoring": "Báo cáo thay đổi SRS - StudySet & Question Authoring",
}


def set_run_fonts(paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = "Calibri"
        if run._element.rPr is not None:
            run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def replace_paragraph(paragraph) -> int:
    text = paragraph.text
    if text not in TEXT_MAP:
        set_run_fonts(paragraph)
        return 0
    replacement = TEXT_MAP[text]
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)
    set_run_fonts(paragraph)
    return 1


def replace_cell(cell) -> int:
    changed = 0
    for paragraph in cell.paragraphs:
        changed += replace_paragraph(paragraph)
    return changed


def main() -> None:
    doc = Document(DOCX)
    changed = 0
    for paragraph in doc.paragraphs:
        changed += replace_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                changed += replace_cell(cell)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            changed += replace_paragraph(paragraph)
        for paragraph in section.footer.paragraphs:
            changed += replace_paragraph(paragraph)
    doc.save(DOCX)
    print(f"updated={changed}")
    print(DOCX)


if __name__ == "__main__":
    main()
